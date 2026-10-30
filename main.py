
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import mysql.connector
import hashlib
import secrets
from tkinter import simpledialog

# ---------------- DATABASE CONNECTION ----------------
def get_connection():
    return mysql.connector.connect(
        host="localhost",
        user="root",
        password="Umesh@123",
        database="attendance_db"
    )


# ---------------- LOGIN AND USER ROLES ----------------
CURRENT_USER = None


def hash_password(password, salt=None):
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 200000)
    return salt.hex(), digest.hex()


def verify_password(password, salt_hex, hash_hex):
    salt = bytes.fromhex(salt_hex)
    return hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 200000).hex() == hash_hex


def prepare_auth_tables():
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS app_users (
            id INT AUTO_INCREMENT PRIMARY KEY,
            username VARCHAR(60) NOT NULL UNIQUE,
            password_salt VARCHAR(32) NOT NULL,
            password_hash VARCHAR(64) NOT NULL,
            role ENUM('admin','teacher') NOT NULL,
            subject_id INT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            CONSTRAINT fk_user_subject FOREIGN KEY (subject_id) REFERENCES subjects(id)
        ) ENGINE=InnoDB
    """)
    cur.execute("SELECT COUNT(*) FROM app_users WHERE role='admin'")
    admin_count = cur.fetchone()[0]
    conn.commit()
    cur.close()
    conn.close()
    return admin_count


def password_dialog(parent, title, heading, confirm=False):
    result = {"value": None}
    dialog = tk.Toplevel(parent)
    dialog.title(title)
    dialog.geometry("490x380" if confirm else "490x380")
    dialog.resizable(False, False)
    dialog.transient(parent)
    dialog.grab_set()
    tk.Label(dialog, text=heading, font=("Segoe UI", 13, "bold")).pack(pady=(18, 12))
    tk.Label(dialog, text="Username").pack(anchor="w", padx=35)
    username = tk.Entry(dialog, width=34)
    username.pack(padx=35, pady=(3, 9), ipady=4)
    tk.Label(dialog, text="Password").pack(anchor="w", padx=35)
    password = tk.Entry(dialog, width=34, show="*")
    password.pack(padx=35, pady=(3, 9), ipady=4)
    confirm_entry = None
    if confirm:
        tk.Label(dialog, text="Confirm Password").pack(anchor="w", padx=35)
        confirm_entry = tk.Entry(dialog, width=34, show="*")
        confirm_entry.pack(padx=35, pady=(3, 8), ipady=4)

    def submit():
        u, pw = username.get().strip(), password.get()
        if not u or not pw:
            messagebox.showwarning("Missing details", "Enter username and password.", parent=dialog)
            return
        if len(pw) < 8:
            messagebox.showwarning("Weak password", "Password must be at least 8 characters.", parent=dialog)
            return
        if confirm and pw != confirm_entry.get():
            messagebox.showerror("Password mismatch", "Both passwords must match.", parent=dialog)
            return
        result["value"] = (u, pw)
        dialog.destroy()

    buttons = tk.Frame(dialog)
    buttons.pack(pady=12)
    tk.Button(buttons, text="Continue", command=submit, width=12).pack(side="left", padx=5)
    tk.Button(buttons, text="Cancel", command=dialog.destroy, width=12).pack(side="left", padx=5)
    dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
    username.focus_set()
    parent.wait_window(dialog)
    return result["value"]


def login_user():
    global CURRENT_USER
    try:
        admin_count = prepare_auth_tables()
    except mysql.connector.Error as error:
        messagebox.showerror("Database Error", f"Could not prepare login tables.\n\n{error}")
        return False

    if admin_count == 0:
        setup = password_dialog(root, "Create Admin", "Create the first Admin account", confirm=True)
        if not setup:
            return False
        username, password = setup
        salt, digest = hash_password(password)
        try:
            conn = get_connection()
            cur = conn.cursor()
            cur.execute("INSERT INTO app_users (username,password_salt,password_hash,role) VALUES (%s,%s,%s,'admin')",
                        (username, salt, digest))
            conn.commit()
            cur.close()
            conn.close()
            messagebox.showinfo("Admin created", "Admin account created. Please log in.")
        except mysql.connector.Error as error:
            messagebox.showerror("Setup Error", str(error))
            return False

    login = tk.Toplevel(root)
    login.title("Login - Student Attendance")
    login.geometry("390x260")
    login.resizable(False, False)
    login.transient(root)
    login.grab_set()
    tk.Label(login, text="STUDENT ATTENDANCE", font=("Segoe UI", 16, "bold"), fg="navy").pack(pady=(25, 18))
    tk.Label(login, text="Username").pack(anchor="w", padx=45)
    user_entry = tk.Entry(login, width=35)
    user_entry.pack(padx=45, pady=(3, 10), ipady=5)
    tk.Label(login, text="Password").pack(anchor="w", padx=45)
    pass_entry = tk.Entry(login, width=35, show="*")
    pass_entry.pack(padx=45, pady=(3, 12), ipady=5)
    login_result = {"ok": False}

    def do_login():
        global CURRENT_USER
        username, password = user_entry.get().strip(), pass_entry.get()
        if not username or not password:
            messagebox.showwarning("Missing details", "Enter username and password.", parent=login)
            return
        try:
            conn = get_connection()
            cur = conn.cursor(dictionary=True)
            cur.execute("SELECT id, username, password_salt, password_hash, role, subject_id FROM app_users WHERE username=%s", (username,))
            account = cur.fetchone()
            cur.close()
            conn.close()
            if account and verify_password(password, account["password_salt"], account["password_hash"]):
                CURRENT_USER = account
                login_result["ok"] = True
                login.destroy()
            else:
                messagebox.showerror("Login failed", "Incorrect username or password.", parent=login)
        except mysql.connector.Error as error:
            messagebox.showerror("Database Error", str(error), parent=login)

    tk.Button(login, text="LOGIN", command=do_login, bg="navy", fg="white", width=20, pady=8).pack()
    login.protocol("WM_DELETE_WINDOW", login.destroy)
    user_entry.focus_set()
    root.wait_window(login)
    return login_result["ok"]


# ---------------- MAIN WINDOW ----------------
root = tk.Tk()
root.withdraw()
root.title("Student Attendance Management System")
root.geometry("1250x850")
root.state("zoomed")
root.minsize(950, 600)
root.configure(bg="white")

selected_id = None

# ---------------- COLORS & STYLES ----------------
BG = "white"
NAVY = "navy"
BLUE = "blue"
WHITE = "white"
TEXT = "black"
MUTED = "gray"
GREEN = "green"
RED = "red"

style = ttk.Style()
style.theme_use("clam")

style.configure(
    "Treeview",
    background=WHITE,
    foreground=TEXT,
    rowheight=34,
    fieldbackground=WHITE,
    font=("Segoe UI", 10),
    borderwidth=0
)
style.configure(
    "Treeview.Heading",
    background="lightblue",
    foreground=NAVY,
    font=("Segoe UI", 10, "bold"),
    padding=10
)
style.map("Treeview", background=[("selected", "light blue")],
          foreground=[("selected", NAVY)])

style.configure("TCombobox", padding=7)


# ---------------- HEADER ----------------
header = tk.Frame(root, bg=NAVY, height=90)
header.pack(fill="x")
header.pack_propagate(False)
tk.Label(
    header,
    text="STUDENT ATTENDANCE",
    bg=NAVY,
    fg=WHITE,
    font=("Segoe UI", 21, "bold")
).pack(anchor="w", padx=25, pady=(15, 0))

tk.Label(
    header,
    text="Management System  |  Student Records",
    bg=NAVY,
    fg="lightgray",
    font=("Segoe UI", 10)
).pack(anchor="w", padx=27, pady=3)


# ---------------- DASHBOARD CARDS ----------------
cards = tk.Frame(root, bg=BG)
cards.pack(fill="x", padx=22, pady=18)

def create_card(parent, title, color):
    card = tk.Frame(parent, bg=WHITE, padx=18, pady=12,
                    highlightbackground="gainsboro",
                    highlightthickness=1)
    card.pack(side="left", fill="x", expand=True, padx=6)

    tk.Label(card, text=title, bg=WHITE, fg=MUTED,
             font=("Segoe UI", 10)).pack(anchor="w")

    value = tk.Label(card, text="0", bg=WHITE, fg=color,
                     font=("Segoe UI", 23, "bold"))
    value.pack(anchor="w", pady=(4, 0))
    return value

total_value = create_card(cards, "TOTAL STUDENTS", BLUE)
department_value = create_card(cards, "DEPARTMENTS", GREEN)
semester_value = create_card(cards, "SEMESTERS", "purple")


# ---------------- MAIN CONTENT ----------------
content = tk.Frame(root, bg=BG)
content.pack(fill="both", expand=True, padx=22, pady=(0, 18))

# Student form
form = tk.Frame(content, bg=WHITE, padx=18, pady=16,
                highlightbackground="gainsboro",
                highlightthickness=1)
form.pack(side="left", fill="y", padx=(0, 14))

tk.Label(form, text="Student Details", bg=WHITE, fg=NAVY,
         font=("Segoe UI", 15, "bold")).pack(anchor="w", pady=(0, 14))


def make_label(text):
    tk.Label(form, text=text, bg=WHITE, fg=TEXT,
             font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(7, 4))


def make_entry():
    entry = tk.Entry(
        form, font=("Segoe UI", 10), width=27,
        relief="solid", bd=1
    )
    entry.pack(fill="x", ipady=6)
    return entry


make_label("Roll Number")
roll_entry = make_entry()

make_label("Student Name")
name_entry = make_entry()

make_label("Department")
department_entry = make_entry()

make_label("Semester")
semester_combo = ttk.Combobox(
    form, values=("1", "2", "3", "4", "5", "6", "7", "8"),
    state="readonly", width=25
)
semester_combo.pack(fill="x", ipady=4)

make_label("Phone Number")
phone_entry = make_entry()


def make_button(parent, text, color, command):
    return tk.Button(
        parent, text=text, command=command,
        bg=color, fg=WHITE, activebackground=color,
        activeforeground=WHITE, relief="flat",
        cursor="hand2", font=("Segoe UI", 10, "bold"),
        padx=8, pady=9
    )


# ---------------- TABLE AREA ----------------
table_panel = tk.Frame(content, bg=WHITE, padx=15, pady=15,
                       highlightbackground="gainsboro",
                       highlightthickness=1)
table_panel.pack(side="left", fill="both", expand=True)

tk.Label(table_panel, text="Student Records", bg=WHITE,
         fg=NAVY, font=("Segoe UI", 15, "bold")).pack(anchor="w")

search_row = tk.Frame(table_panel, bg=WHITE)
search_row.pack(fill="x", pady=12)

search_entry = tk.Entry(
    search_row, font=("Segoe UI", 10),
    relief="solid", bd=1
)
search_entry.pack(side="left", fill="x", expand=True, ipady=8)

tree_frame = tk.Frame(table_panel, bg=WHITE)
tree_frame.pack(fill="both", expand=True)

columns = ("id", "roll", "name", "department", "semester", "phone")
tree = ttk.Treeview(
    tree_frame, columns=columns,
    show="headings", selectmode="browse"
)

headings = {
    "id": ("ID", 45),
    "roll": ("Roll No.", 85),
    "name": ("Student Name", 145),
    "department": ("Department", 115),
    "semester": ("Sem.", 55),
    "phone": ("Phone", 105)
}

for col, (heading, width) in headings.items():
    tree.heading(col, text=heading)
    tree.column(col, width=width, minwidth=40, anchor="center")

scrollbar = ttk.Scrollbar(tree_frame, orient="vertical",
                          command=tree.yview)
tree.configure(yscrollcommand=scrollbar.set)
tree.pack(side="left", fill="both", expand=True)
scrollbar.pack(side="right", fill="y")


# ---------------- FUNCTIONS ----------------
def clear_form():
    global selected_id
    selected_id = None

    for entry in (roll_entry, name_entry,
                  department_entry, phone_entry):
        entry.delete(0, tk.END)

    semester_combo.set("")
    tree.selection_remove(tree.selection())


def refresh_students(*args):
    try:
        conn = get_connection()
        cur = conn.cursor()

        cur.execute("""
            SELECT id, roll_no, name, department, semester, phone
            FROM students
            ORDER BY id DESC
        """)
        rows = cur.fetchall()

        for item in tree.get_children():
            tree.delete(item)

        query = search_entry.get().strip().lower()

        for row in rows:
            searchable = " ".join(
                str(value or "") for value in row
            ).lower()

            if query and query not in searchable:
                continue

            tree.insert("", "end", values=row)

        cur.execute("SELECT COUNT(*) FROM students")
        total_value.config(text=str(cur.fetchone()[0]))

        cur.execute("""
            SELECT COUNT(DISTINCT department)
            FROM students
            WHERE department IS NOT NULL
              AND department <> ''
        """)
        department_value.config(text=str(cur.fetchone()[0]))

        cur.execute("""
            SELECT COUNT(DISTINCT semester)
            FROM students
            WHERE semester IS NOT NULL
        """)
        semester_value.config(text=str(cur.fetchone()[0]))

        cur.close()
        conn.close()

    except mysql.connector.Error as error:
        messagebox.showerror(
            "Database Error",
            f"Could not load student records.\n\n{error}"
        )


def save_student():
    global selected_id

    roll = roll_entry.get().strip()
    name = name_entry.get().strip()
    department = department_entry.get().strip()
    semester = semester_combo.get()
    phone = phone_entry.get().strip()

    if not roll or not name:
        messagebox.showwarning(
            "Missing Details",
            "Roll Number and Student Name are required."
        )
        return

    if phone and not phone.isdigit():
        messagebox.showwarning(
            "Invalid Phone",
            "Phone Number must contain digits only."
        )
        return

    try:
        conn = get_connection()
        cur = conn.cursor()

        values = (
            roll, name, department or None,
            int(semester) if semester else None,
            phone or None
        )

        if selected_id is None:
            cur.execute("""
                INSERT INTO students
                (roll_no, name, department, semester, phone)
                VALUES (%s, %s, %s, %s, %s)
            """, values)
            message = "Student added successfully!"
        else:
            cur.execute("""
                UPDATE students
                SET roll_no=%s, name=%s, department=%s,
                    semester=%s, phone=%s
                WHERE id=%s
            """, values + (selected_id,))
            message = "Student updated successfully!"

        conn.commit()
        cur.close()
        conn.close()

        clear_form()
        refresh_students()
        messagebox.showinfo("Success", message)

    except mysql.connector.IntegrityError:
        messagebox.showerror(
            "Duplicate Roll Number",
            "This Roll Number already exists."
        )
    except mysql.connector.Error as error:
        messagebox.showerror("Database Error", str(error))


def select_student(event=None):
    global selected_id

    selection = tree.selection()
    if not selection:
        return

    values = tree.item(selection[0], "values")
    selected_id = int(values[0])

    clear_entries = (
        (roll_entry, values[1]),
        (name_entry, values[2]),
        (department_entry, values[3]),
        (phone_entry, values[5])
    )

    for entry, value in clear_entries:
        entry.delete(0, tk.END)
        entry.insert(0, value or "")

    semester_combo.set(
        str(values[4]) if values[4] else ""
    )


def delete_student():
    selection = tree.selection()

    if not selection:
        messagebox.showwarning(
            "Select Student",
            "Select a student from the table first."
        )
        return

    values = tree.item(selection[0], "values")
    student_id = int(values[0])

    confirm = messagebox.askyesno(
        "Confirm Delete",
        f"Delete student '{values[2]}'?"
    )

    if not confirm:
        return

    try:
        conn = get_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM students WHERE id=%s",
                    (student_id,))
        conn.commit()
        cur.close()
        conn.close()

        clear_form()
        refresh_students()
        messagebox.showinfo("Success", "Student deleted.")

    except mysql.connector.Error as error:
        messagebox.showerror("Database Error", str(error))

# ---------------- ATTENDANCE MANAGEMENT ----------------
from datetime import date


def open_attendance_window():
    win = tk.Toplevel(root)
    win.title("Attendance Management")
    win.geometry("1100x750")
    win.minsize(900, 650)
    win.state("zoomed")
    win.configure(bg=BG)
    win.transient(root)

    tk.Label(
        win,
        text="ATTENDANCE MANAGEMENT",
        bg=NAVY, fg=WHITE,
        font=("Segoe UI", 18, "bold"),
        padx=20, pady=15
    ).pack(fill="x")

    top = tk.Frame(win, bg=WHITE, padx=15, pady=12)
    top.pack(fill="x", padx=18, pady=15)

    tk.Label(
        top, text="Attendance Date:",
        bg=WHITE, fg=TEXT,
        font=("Segoe UI", 10, "bold")
    ).pack(side="left")

    date_entry = tk.Entry(top, font=("Segoe UI", 11), width=14)
    date_entry.pack(side="left", padx=10, ipady=5)
    date_entry.insert(0, date.today().isoformat())

    tk.Label(
        top, text="Format: YYYY-MM-DD",
        bg=WHITE, fg=MUTED,
        font=("Segoe UI", 9)
    ).pack(side="left")

    tk.Label(top, text="Subject:", bg=WHITE, fg=TEXT,
             font=("Segoe UI", 10, "bold")).pack(side="left", padx=(25, 6))
    subject_var = tk.StringVar()
    subject_combo = ttk.Combobox(top, textvariable=subject_var,
                                 state="readonly", width=24)
    subject_combo.pack(side="left")
    subject_lookup = {}
    try:
        conn = get_connection()
        cur = conn.cursor()
        if CURRENT_USER["role"] == "teacher":
            cur.execute("SELECT id, subject_name FROM subjects WHERE id=%s", (CURRENT_USER["subject_id"],))
        else:
            cur.execute("SELECT id, subject_name FROM subjects WHERE subject_name <> 'Previous Records' ORDER BY subject_name")
        subject_rows = cur.fetchall()
        cur.close()
        conn.close()
        subject_lookup = {name: sid for sid, name in subject_rows}
        subject_combo["values"] = list(subject_lookup.keys())
        if subject_rows:
            subject_combo.current(0)
        if CURRENT_USER["role"] == "teacher":
            subject_combo.configure(state="disabled")
        if not subject_rows:
            messagebox.showerror("No Subject", "No subject is assigned/available.", parent=win)
    except mysql.connector.Error as error:
        messagebox.showerror("Database Error", str(error), parent=win)

    body = tk.Frame(win, bg=WHITE, padx=15, pady=12)
    body.pack(fill="both", expand=True, padx=18, pady=(0, 15))

    cols = ("id", "roll", "name", "department", "status")
    attendance_tree = ttk.Treeview(
        body, columns=cols, show="headings",
        selectmode="browse"
    )

    headers = {
        "id": ("ID", 50),
        "roll": ("Roll No.", 110),
        "name": ("Student Name", 220),
        "department": ("Department", 160),
        "status": ("Attendance Status", 160)
    }

    for col, (heading, width) in headers.items():
        attendance_tree.heading(col, text=heading)
        attendance_tree.column(col, width=width, anchor="center")

    attendance_tree.pack(fill="both", expand=True)

    selected_status = tk.StringVar(value="Present")

    controls = tk.Frame(body, bg=WHITE)
    controls.pack(fill="x", pady=(12, 0))

    tk.Label(
        controls, text="Mark selected student:",
        bg=WHITE, fg=TEXT,
        font=("Segoe UI", 10, "bold")
    ).pack(side="left", padx=(0, 10))

    ttk.Combobox(
        controls, textvariable=selected_status,
        values=("Present", "Absent"),
        state="readonly", width=12
    ).pack(side="left")

    def load_attendance():
        try:
            chosen_date = date.fromisoformat(date_entry.get().strip())

            conn = get_connection()
            cur = conn.cursor()

            subject_id = subject_lookup.get(subject_var.get())
            if not subject_id:
                messagebox.showwarning("Select Subject", "Select a subject first.", parent=win)
                return
            cur.execute("""
                SELECT s.id, s.roll_no, s.name, s.department,
                       COALESCE(a.status, 'Not Marked')
                FROM students s
                LEFT JOIN attendance a
                  ON s.id = a.student_id
                 AND a.attendance_date = %s
                 AND a.subject_id = %s
                ORDER BY s.roll_no
            """, (chosen_date, subject_id))

            rows = cur.fetchall()

            for item in attendance_tree.get_children():
                attendance_tree.delete(item)

            for row in rows:
                attendance_tree.insert("", "end", values=row)

            cur.close()
            conn.close()

        except ValueError:
            messagebox.showerror(
                "Invalid Date",
                "Enter the date in YYYY-MM-DD format.",
                parent=win
            )
        except mysql.connector.Error as error:
            messagebox.showerror(
                "Database Error", str(error), parent=win
            )

    def mark_selected():
        selection = attendance_tree.selection()

        if not selection:
            messagebox.showwarning(
                "Select Student",
                "Please select a student first.",
                parent=win
            )
            return

        try:
            chosen_date = date.fromisoformat(date_entry.get().strip())
            student = attendance_tree.item(
                selection[0], "values"
            )
            student_id = int(student[0])
            status = selected_status.get()
            subject_id = subject_lookup.get(subject_var.get())
            if not subject_id:
                messagebox.showwarning("Select Subject", "Select a subject first.", parent=win)
                return

            conn = get_connection()
            cur = conn.cursor()

            cur.execute("""
                INSERT INTO attendance
                    (student_id, subject_id, attendance_date, status)
                VALUES (%s, %s, %s, %s)
                ON DUPLICATE KEY UPDATE status = %s
            """, (student_id, subject_id, chosen_date, status, status))

            conn.commit()
            cur.close()
            conn.close()

            load_attendance()
            messagebox.showinfo(
                "Saved",
                f"{student[2]} marked {status}.",
                parent=win
            )

        except ValueError:
            messagebox.showerror(
                "Invalid Date",
                "Enter the date in YYYY-MM-DD format.",
                parent=win
            )
        except mysql.connector.Error as error:
            messagebox.showerror(
                "Database Error", str(error), parent=win
            )

    actions = tk.Frame(win, bg=BG)
    actions.pack(fill="x", padx=18, pady=(0, 15))

    make_button(
        actions, "Load / Refresh", BLUE,
        load_attendance
    ).pack(side="left", padx=(0, 8))

    make_button(
        actions, "Save Attendance", GREEN,
        mark_selected
    ).pack(side="left")

    subject_combo.bind("<<ComboboxSelected>>", lambda event: load_attendance())
    load_attendance()


# ---------------- ADMIN: SUBJECT AND TEACHER MANAGEMENT ----------------
def open_teacher_management():
    if CURRENT_USER["role"] != "admin":
        messagebox.showerror("Access denied", "Only Admin can manage teachers.")
        return
    win = tk.Toplevel(root)
    win.title("Manage Teachers and Subjects")
    win.geometry("650x500")
    win.configure(bg=BG)
    tk.Label(win, text="TEACHER & SUBJECT MANAGEMENT", bg=NAVY, fg=WHITE,
             font=("Segoe UI", 16, "bold"), pady=14).pack(fill="x")

    subject_frame = tk.LabelFrame(win, text="Add Subject", padx=12, pady=12)
    subject_frame.pack(fill="x", padx=18, pady=15)
    subject_entry = tk.Entry(subject_frame, width=35)
    subject_entry.pack(side="left", padx=(0, 10), ipady=5)

    def add_subject():
        name = subject_entry.get().strip()
        if not name:
            messagebox.showwarning("Missing", "Enter a subject name.", parent=win)
            return
        try:
            conn = get_connection(); cur = conn.cursor()
            cur.execute("INSERT INTO subjects (subject_name) VALUES (%s)", (name,))
            conn.commit(); cur.close(); conn.close()
            subject_entry.delete(0, tk.END)
            refresh_subjects()
            messagebox.showinfo("Success", "Subject added.", parent=win)
        except mysql.connector.IntegrityError:
            messagebox.showerror("Duplicate", "This subject already exists.", parent=win)
        except mysql.connector.Error as error:
            messagebox.showerror("Database Error", str(error), parent=win)

    tk.Button(subject_frame, text="Add Subject", command=add_subject).pack(side="left")

    teacher_frame = tk.LabelFrame(win, text="Create Teacher Account", padx=12, pady=12)
    teacher_frame.pack(fill="x", padx=18, pady=8)
    tk.Label(teacher_frame, text="Teacher username").grid(row=0, column=0, sticky="w", pady=5)
    username_entry = tk.Entry(teacher_frame, width=30); username_entry.grid(row=0, column=1, padx=8, pady=5)
    tk.Label(teacher_frame, text="Temporary password (min 8 chars)").grid(row=1, column=0, sticky="w", pady=5)
    password_entry = tk.Entry(teacher_frame, width=30, show="*"); password_entry.grid(row=1, column=1, padx=8, pady=5)
    tk.Label(teacher_frame, text="Assigned subject").grid(row=2, column=0, sticky="w", pady=5)
    subject_var = tk.StringVar()
    teacher_subject_combo = ttk.Combobox(teacher_frame, textvariable=subject_var, state="readonly", width=27)
    teacher_subject_combo.grid(row=2, column=1, padx=8, pady=5)

    def refresh_subjects():
        try:
            conn = get_connection(); cur = conn.cursor()
            cur.execute("SELECT id, subject_name FROM subjects WHERE subject_name <> 'Previous Records' ORDER BY subject_name")
            rows = cur.fetchall(); cur.close(); conn.close()
            win.subject_map = {name: sid for sid, name in rows}
            teacher_subject_combo["values"] = list(win.subject_map.keys())
            if rows and not subject_var.get(): teacher_subject_combo.current(0)
        except mysql.connector.Error as error:
            messagebox.showerror("Database Error", str(error), parent=win)

    def create_teacher():
        username, password, subject_name = username_entry.get().strip(), password_entry.get(), subject_var.get()
        if not username or not password or not subject_name:
            messagebox.showwarning("Missing", "Enter username, password and assigned subject.", parent=win)
            return
        if len(password) < 8:
            messagebox.showwarning("Weak password", "Password must be at least 8 characters.", parent=win)
            return
        salt, digest = hash_password(password)
        try:
            conn = get_connection(); cur = conn.cursor()
            cur.execute("INSERT INTO app_users (username,password_salt,password_hash,role,subject_id) VALUES (%s,%s,%s,'teacher',%s)",
                        (username, salt, digest, win.subject_map[subject_name]))
            conn.commit(); cur.close(); conn.close()
            username_entry.delete(0, tk.END); password_entry.delete(0, tk.END)
            messagebox.showinfo("Success", "Teacher account created.", parent=win)
        except mysql.connector.IntegrityError:
            messagebox.showerror("Duplicate username", "This username already exists.", parent=win)
        except mysql.connector.Error as error:
            messagebox.showerror("Database Error", str(error), parent=win)

    tk.Button(teacher_frame, text="Create Teacher", command=create_teacher,
              bg=GREEN, fg=WHITE, padx=12, pady=7).grid(row=3, column=1, sticky="e", pady=10)
    refresh_subjects()


# ---------------- ATTENDANCE REPORTS ----------------
def open_attendance_reports():
    report_win = tk.Toplevel(root)
    report_win.title("Attendance Reports")
    report_win.geometry("1150x700")
    report_win.minsize(900, 600)
    report_win.configure(bg=BG)
    report_win.transient(root)

    tk.Label(
        report_win, text="ATTENDANCE REPORTS",
        bg=NAVY, fg=WHITE, font=("Segoe UI", 18, "bold"),
        padx=20, pady=15
    ).pack(fill="x")

    filters = tk.Frame(report_win, bg=WHITE, padx=15, pady=12)
    filters.pack(fill="x", padx=18, pady=15)

    tk.Label(filters, text="Report Period:", bg=WHITE, fg=TEXT,
             font=("Segoe UI", 10, "bold")).pack(side="left")
    period_combo = ttk.Combobox(
        filters, values=("All Time", "Monthly"),
        state="readonly", width=12
    )
    period_combo.set("All Time")
    period_combo.pack(side="left", padx=(8, 18))

    tk.Label(filters, text="Month:", bg=WHITE, fg=TEXT,
             font=("Segoe UI", 10, "bold")).pack(side="left")
    month_combo = ttk.Combobox(
        filters, values=tuple(f"{m:02d}" for m in range(1, 13)),
        state="readonly", width=5
    )
    month_combo.set(f"{date.today().month:02d}")
    month_combo.pack(side="left", padx=(8, 15))

    tk.Label(filters, text="Year:", bg=WHITE, fg=TEXT,
             font=("Segoe UI", 10, "bold")).pack(side="left")
    year_combo = ttk.Combobox(
        filters,
        values=tuple(str(y) for y in range(date.today().year - 10, date.today().year + 2)),
        state="readonly", width=7
    )
    year_combo.set(str(date.today().year))
    year_combo.pack(side="left", padx=(8, 15))

    report_frame = tk.Frame(report_win, bg=WHITE, padx=12, pady=12)
    report_frame.pack(fill="both", expand=True, padx=18, pady=(0, 12))

    report_columns = ("roll", "name", "department", "present", "absent", "marked", "percentage")
    report_tree = ttk.Treeview(
        report_frame, columns=report_columns, show="headings", selectmode="browse"
    )
    report_headers = {
        "roll": ("Roll No.", 100),
        "name": ("Student Name", 200),
        "department": ("Department", 180),
        "present": ("Present Days", 110),
        "absent": ("Absent Days", 110),
        "marked": ("Marked Days", 110),
        "percentage": ("Attendance %", 120),
    }
    for col, (heading, width) in report_headers.items():
        report_tree.heading(col, text=heading)
        report_tree.column(col, width=width, minwidth=70, anchor="center")

    report_scroll = ttk.Scrollbar(report_frame, orient="vertical", command=report_tree.yview)
    report_tree.configure(yscrollcommand=report_scroll.set)
    report_tree.pack(side="left", fill="both", expand=True)
    report_scroll.pack(side="right", fill="y")

    summary_label = tk.Label(
        report_win, text="", bg=BG, fg=NAVY,
        font=("Segoe UI", 10, "bold")
    )
    summary_label.pack(anchor="w", padx=20, pady=(0, 10))

    def fetch_report_rows():
        conn = get_connection()
        cur = conn.cursor()
        teacher_filter = " AND a.subject_id = %s" if CURRENT_USER["role"] == "teacher" else ""
        teacher_params = (CURRENT_USER["subject_id"],) if CURRENT_USER["role"] == "teacher" else ()
        if period_combo.get() == "Monthly":
            month = month_combo.get()
            year = year_combo.get()
            if not month or not year:
                cur.close()
                conn.close()
                raise ValueError("Please select both month and year.")
            cur.execute("""
                SELECT s.roll_no, s.name, COALESCE(s.department, ''),
                       COALESCE(SUM(a.status = 'Present'), 0) AS present_days,
                       COALESCE(SUM(a.status = 'Absent'), 0) AS absent_days
                FROM students s
                LEFT JOIN attendance a
                  ON a.student_id = s.id
                 AND YEAR(a.attendance_date) = %s
                 AND MONTH(a.attendance_date) = %s
                  {teacher_filter}
                GROUP BY s.id, s.roll_no, s.name, s.department
                ORDER BY s.roll_no
            """.format(teacher_filter=teacher_filter), (int(year), int(month)) + teacher_params)
        else:
            cur.execute("""
                SELECT s.roll_no, s.name, COALESCE(s.department, ''),
                       COALESCE(SUM(a.status = 'Present'), 0) AS present_days,
                       COALESCE(SUM(a.status = 'Absent'), 0) AS absent_days
                FROM students s
                LEFT JOIN attendance a ON a.student_id = s.id
                  {teacher_filter}
                GROUP BY s.id, s.roll_no, s.name, s.department
                ORDER BY s.roll_no
            """.format(teacher_filter=teacher_filter), teacher_params)
        rows = cur.fetchall()
        cur.close()
        conn.close()

        output = []
        for roll, name, department, present, absent in rows:
            present = int(present or 0)
            absent = int(absent or 0)
            marked = present + absent
            percentage = (present / marked * 100) if marked else 0.0
            output.append((roll, name, department, present, absent, marked, percentage))
        return output

    def load_report():
        try:            rows = fetch_report_rows()
            for item in report_tree.get_children():
                report_tree.delete(item)
            for row in rows:
                report_tree.insert("", "end", values=(
                    row[0], row[1], row[2], row[3], row[4], row[5], f"{row[6]:.2f}%"
                ))
            total_present = sum(row[3] for row in rows)
            total_absent = sum(row[4] for row in rows)
            total_marked = total_present + total_absent
            overall = total_present / total_marked * 100 if total_marked else 0.0
            period_text = "all recorded dates" if period_combo.get() == "All Time" else f"{month_combo.get()}/{year_combo.get()}"
            summary_label.config(
                text=f"Students: {len(rows)}    |    Present: {total_present}    |    Absent: {total_absent}    |    Overall attendance: {overall:.2f}%    |    Period: {period_text}"
            )
        except ValueError as error:
            messagebox.showwarning("Select Period", str(error), parent=report_win)
        except mysql.connector.Error as error:
            messagebox.showerror("Database Error", str(error), parent=report_win)

    def export_excel():
        try:
            rows = fetch_report_rows()
            if not rows:
                messagebox.showinfo("No Data", "There are no student records to export.", parent=report_win)
                return
            filename = filedialog.asksaveasfilename(
                parent=report_win,
                title="Save Attendance Report",
                defaultextension=".xlsx",
                filetypes=(("Excel Workbook", "*.xlsx"),),
                initialfile=("attendance_all_time.xlsx" if period_combo.get() == "All Time"
                             else f"attendance_{year_combo.get()}_{month_combo.get()}.xlsx")
            )
            if not filename:
                return

            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Attendance Report"
            sheet.append(["Roll No.", "Student Name", "Department", "Present Days",
                          "Absent Days", "Marked Days", "Attendance Percentage"])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="172554")
                cell.alignment = Alignment(horizontal="center")

            for row in rows:
                sheet.append([row[0], row[1], row[2], row[3], row[4], row[5], round(row[6], 2)])
                sheet.cell(sheet.max_row, 7).number_format = '0.00"%"'

            for column_cells in sheet.columns:
                max_length = max(len(str(cell.value or "")) for cell in column_cells)
                sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_length + 3, 32)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            workbook.save(filename)
            messagebox.showinfo("Export Successful", f"Excel report saved to:\n{filename}", parent=report_win)
        except ValueError as error:
            messagebox.showwarning("Select Period", str(error), parent=report_win)
        except mysql.connector.Error as error:
            messagebox.showerror("Database Error", str(error), parent=report_win)
        except Exception as error:
            messagebox.showerror("Export Error", str(error), parent=report_win)

    period_combo.bind("<<ComboboxSelected>>", lambda event: load_report())
    month_combo.bind("<<ComboboxSelected>>", lambda event: load_report())
    year_combo.bind("<<ComboboxSelected>>", lambda event: load_report())

    buttons = tk.Frame(report_win, bg=BG)
    buttons.pack(fill="x", padx=18, pady=(0, 15))
    make_button(buttons, "Refresh Report", BLUE, load_report).pack(side="left", padx=(0, 8))
    make_button(buttons, "Export to Excel", GREEN, export_excel).pack(side="left")
    load_report()




# ---------------- BUTTONS ----------------
make_button(form, "Save Student", BLUE,
            save_student).pack(fill="x", pady=(18, 7))

make_button(form, "Clear Form", "gray",
            clear_form).pack(fill="x", pady=4)

make_button(search_row, "Search", BLUE,
            refresh_students).pack(side="left", padx=(8, 0))

make_button(search_row, "Show All", "gray",
            lambda: (search_entry.delete(0, tk.END),
                     refresh_students())).pack(side="left", padx=(6, 0))

tree.bind("<<TreeviewSelect>>", select_student)

action_row = tk.Frame(table_panel, bg=WHITE)
action_row.pack(fill="x", pady=(12, 0))

make_button(action_row, "Update Selected", GREEN,
            save_student).pack(side="left", padx=(0, 8))

make_button(action_row, "Delete Selected", RED,
            delete_student).pack(side="left")

make_button(action_row, "Manage Attendance", "purple",
            open_attendance_window).pack(side="left", padx=8)

make_button(action_row, "Attendance Reports", "teal",
            open_attendance_reports).pack(side="left", padx=8)

make_button(action_row, "Refresh", BLUE,
            refresh_students).pack(side="right")


# ---------------- START APP ----------------
if login_user():
    root.deiconify()
    if CURRENT_USER["role"] == "admin":
        refresh_students()
    root.mainloop()
else:
    root.destroy()